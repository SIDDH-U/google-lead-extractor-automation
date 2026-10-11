import os
import sys
from pathlib import Path
import openpyxl

# Ensure UTF-8 stdout on Windows
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from cid_utils import extract_cid_from_url, hex_to_decimal_cid, get_short_cid_url
from db import DatabaseManager
from exporter import export_leads_to_excel
from geo import get_city_bounding_box, generate_grid_points, calculate_haversine_distance_km


def test_cid_utils():
    print("[*] Testing CID Utils...")
    # Test hex to decimal
    hex_sample = "0xd134e199e405a163"
    dec_sample = hex_to_decimal_cid(hex_sample)
    assert dec_sample is not None, "Failed to convert hex to decimal"
    assert dec_sample.isdigit(), "Decimal CID must be digits"

    # Test URL parsing with !1s0x...:0x...
    test_url = "https://www.google.com/maps/place/Empire+State+Building/@40.7484405,-73.9856644,17z/data=!4m6!3m5!1s0x89c259a9b3117469:0xd134e199e405a163!8m2!3d40.7484405!4d-73.9856644"
    hex_cid, dec_cid = extract_cid_from_url(test_url)
    assert hex_cid == "d134e199e405a163", f"Expected hex cid 'd134e199e405a163', got '{hex_cid}'"
    assert dec_cid == "15074921903787712867", f"Expected decimal cid '15074921903787712867', got '{dec_cid}'"

    short_link = get_short_cid_url(test_url)
    assert short_link == "https://maps.google.com/?cid=15074921903787712867", f"Unexpected short link: {short_link}"
    print("  ✓ CID extraction & short URL conversion passed!")


def test_geo_grid():
    print("[*] Testing Geo & Grid Generator...")
    # Mock bbox for test
    mock_bbox = {
        "city": "Austin, TX",
        "display_name": "Austin, Travis County, Texas, United States",
        "lat_min": 30.15,
        "lat_max": 30.35,
        "lon_min": -97.85,
        "lon_max": -97.65,
        "center_lat": 30.25,
        "center_lon": -97.75,
    }
    points = generate_grid_points(mock_bbox, step_km=5.0)
    assert len(points) > 0, "Should generate at least 1 grid point"
    print(f"  ✓ Generated {len(points)} grid points for 5km step.")

    # Distance check
    dist = calculate_haversine_distance_km(30.15, -97.85, 30.35, -97.65)
    assert dist > 10, "Distance should be > 10 km"
    print(f"  ✓ Haversine distance calculated: {dist:.2f} km")


def test_db_manager(tmp_path: Path):
    print("[*] Testing Database Manager & JSON Sync...")
    test_db_file = tmp_path / "test_leads.db"
    test_json_file = tmp_path / "test_leads.json"

    if test_db_file.exists():
        test_db_file.unlink()
    if test_json_file.exists():
        test_json_file.unlink()

    db = DatabaseManager(test_db_file, test_json_file)
    
    lead_1 = {
        "cid": "15074811776964174179",
        "business_name": "Austin Plumbing Pros",
        "phone": "+1 512-555-0199",
        "category": "Plumber",
        "address": "123 Congress Ave, Austin, TX 78701",
        "website": "https://austinplumbingpros.example.com",
        "google_maps_url": "https://www.google.com/maps/place/data=!1s0x89c259a9b3117469:0xd134e199e405a163",
        "short_cid_url": "https://maps.google.com/?cid=15074811776964174179",
        "rating": "4.9",
        "reviews_count": "142",
        "keyword": "plumber",
        "city": "Austin, TX",
    }

    is_new = db.save_lead(lead_1)
    assert is_new is True, "First save should be new"
    assert db.get_leads_count() == 1, "Database should contain 1 lead"
    assert db.is_cid_saved("15074811776964174179") is True

    # Test update / deduplication
    lead_1_updated = dict(lead_1)
    lead_1_updated["phone"] = "+1 512-555-0200"
    is_new_2 = db.save_lead(lead_1_updated)
    assert is_new_2 is False, "Second save of same CID should update, not insert new"
    assert db.get_leads_count() == 1, "Count should remain 1 after update"

    # Verify JSON backup was created
    assert test_json_file.exists(), "JSON backup file must exist after lead save"
    print("  ✓ Database insertion, deduplication & JSON sync passed!")


def test_excel_export(tmp_path: Path):
    print("[*] Testing Excel Export with Dropdowns & Styling...")
    test_excel_file = tmp_path / "test_export.xlsx"

    sample_leads = [
        {
            "business_name": "Austin Plumbing Pros",
            "phone": "+1 512-555-0199",
            "short_cid_url": "https://maps.google.com/?cid=15074811776964174179",
            "google_maps_url": "https://maps.google.com/?cid=15074811776964174179",
            "category": "Plumber",
            "call_status": "",
            "interest": "",
            "send": "",
        },
        {
            "business_name": "Lone Star HVAC & Electrical",
            "phone": "+1 512-555-0844",
            "short_cid_url": "https://maps.google.com/?cid=1234567890123456789",
            "google_maps_url": "https://maps.google.com/?cid=1234567890123456789",
            "category": "HVAC Contractor",
            "call_status": "",
            "interest": "",
            "send": "",
        },
    ]

    out_path = export_leads_to_excel(sample_leads, test_excel_file)
    assert out_path.exists(), "Excel file was not created"

    # Validate Excel content using openpyxl
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active

    # Check Headers
    expected_headers = [
        "Business Name", "Phone", "Google Maps URL", "Business Type",
        "Call Status", "Interest", "Send"
    ]
    actual_headers = [ws.cell(row=1, column=i).value for i in range(1, 8)]
    assert actual_headers == expected_headers, f"Header mismatch: {actual_headers}"

    # Check Row 2 content
    assert ws.cell(row=2, column=1).value == "Austin Plumbing Pros"
    assert ws.cell(row=2, column=2).value == "+1 512-555-0199"
    assert ws.cell(row=2, column=3).hyperlink is not None, "URL must have clickable hyperlink"
    assert ws.cell(row=2, column=4).value == "Plumber"

    # Check Data Validations (Call Status, Interest, Send, Lead Quality)
    data_validations = ws.data_validations.dataValidation
    assert len(data_validations) == 4, f"Expected 4 data validation objects, got {len(data_validations)}"

    # Check Freeze Pane
    assert ws.freeze_panes == "A2", f"Expected freeze pane A2, got {ws.freeze_panes}"
    print("  ✓ Excel export, dropdowns, headers & hyperlinks verified!")


def main():
    tmp_dir = Path("./tmp_test")
    tmp_dir.mkdir(parents=True, exist_ok=True)
    try:
        test_cid_utils()
        test_geo_grid()
        test_db_manager(tmp_dir)
        test_excel_export(tmp_dir)
        print("\n[SUCCESS] All unit tests passed perfectly! 🎉")
    finally:
        # Clean up test artifacts
        import shutil
        shutil.rmtree(tmp_dir, ignore_errors=True)


if __name__ == "__main__":
    main()
